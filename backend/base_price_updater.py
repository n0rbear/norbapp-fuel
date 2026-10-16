"""Daily refresh of the national average base prices (P_base) and FX rates.

Source: European Commission Weekly Oil Bulletin, "prices with taxes" sheet.
The bulletin is published weekly, but the refresh runs every day at
REFRESH_TIME (default 06:00, Europe/Budapest) so a new edition is picked up
the morning after it appears. A refresh also runs at startup, and if the
instance was asleep at the scheduled time, the next request triggers a
catch-up refresh in the background.
"""

from __future__ import annotations

import asyncio
import io
import logging
import os
import re
from datetime import date, datetime, time, timedelta
from html import unescape
from urllib.parse import urljoin
from zoneinfo import ZoneInfo

import httpx
from openpyxl import load_workbook

import pricing_matrix as pm
from fx import fx
from pricing_matrix import Fuel

logger = logging.getLogger(__name__)

BULLETIN_PAGE = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
RETRY_AFTER_FAILURE = timedelta(minutes=30)

COUNTRY_NAMES = {
    "Austria": "AT", "Belgium": "BE", "Bulgaria": "BG", "Croatia": "HR", "Cyprus": "CY",
    "Czechia": "CZ", "Czech Republic": "CZ", "Denmark": "DK", "Estonia": "EE", "Finland": "FI",
    "France": "FR", "Germany": "DE", "Greece": "GR", "Hungary": "HU", "Ireland": "IE",
    "Italy": "IT", "Latvia": "LV", "Lithuania": "LT", "Luxembourg": "LU", "Malta": "MT",
    "Netherlands": "NL", "Poland": "PL", "Portugal": "PT", "Romania": "RO", "Slovakia": "SK",
    "Slovenia": "SI", "Spain": "ES", "Sweden": "SE",
}

_XLSX_LINK = re.compile(r'href="([^"]*prices%20with%20taxes[^"]*\.xlsx[^"]*)"', re.IGNORECASE)


def _tz() -> ZoneInfo:
    return ZoneInfo(os.getenv("REFRESH_TZ", "Europe/Budapest"))


def _refresh_time() -> time:
    hour, minute = os.getenv("REFRESH_TIME", "06:00").split(":")
    return time(int(hour), int(minute))


def last_scheduled_run(now: datetime) -> datetime:
    """Most recent REFRESH_TIME at or before `now` (timezone-aware)."""
    local = now.astimezone(_tz())
    candidate = datetime.combine(local.date(), _refresh_time(), tzinfo=_tz())
    return candidate if candidate <= local else candidate - timedelta(days=1)


def next_scheduled_run(now: datetime) -> datetime:
    return last_scheduled_run(now) + timedelta(days=1)


def find_xlsx_url(page_html: str) -> str:
    match = _XLSX_LINK.search(page_html)
    if not match:
        raise ValueError("Weekly Oil Bulletin 'prices with taxes' link not found")
    return urljoin(BULLETIN_PAGE, unescape(match.group(1)))


def parse_bulletin(content: bytes) -> tuple[dict[str, dict[Fuel, float]], str]:
    """Return ({country: {fuel: EUR/l}}, as_of ISO date) from the bulletin xlsx."""
    sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    prices: dict[str, dict[Fuel, float]] = {}
    as_of: str | None = None
    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        first = row[0]
        if as_of is None and isinstance(first, (datetime, date)):
            as_of = first.date().isoformat() if isinstance(first, datetime) else first.isoformat()
            continue
        code = COUNTRY_NAMES.get(str(first).strip()) if first is not None else None
        if code is None or len(row) < 3:
            continue
        try:
            petrol, diesel = float(row[1]) / 1000, float(row[2]) / 1000  # EUR per 1000 l
        except (TypeError, ValueError):
            continue
        if 0.5 < petrol < 4 and 0.5 < diesel < 4:
            prices[code] = {Fuel.PETROL: round(petrol, 3), Fuel.DIESEL: round(diesel, 3)}
    if len(prices) < 20:
        raise ValueError(f"Bulletin parsed only {len(prices)} countries")
    return prices, as_of or datetime.now(_tz()).date().isoformat()


class BasePriceUpdater:
    def __init__(self) -> None:
        self.enabled = os.getenv("BASE_PRICE_AUTO_REFRESH", "true").lower() not in {"0", "false", "off"}
        self.last_success: datetime | None = None
        self.last_attempt: datetime | None = None
        self.last_error: str | None = None
        self.fx_error: str | None = None
        self.source_url: str | None = None
        self._lock = asyncio.Lock()
        self._task: asyncio.Task | None = None
        self._catch_up: asyncio.Task | None = None

    async def refresh(self, http: httpx.AsyncClient) -> bool:
        async with self._lock:
            self.last_attempt = datetime.now(_tz())
            fx_ok = await self._refresh_fx(http)
            return await self._refresh_prices(http) and fx_ok

    async def _refresh_fx(self, http: httpx.AsyncClient) -> bool:
        try:
            await fx.refresh(http)
        except (httpx.HTTPError, ValueError) as exc:
            self.fx_error = f"{exc.__class__.__name__}: {exc}"[:300]
            logger.warning("FX refresh failed: %s", self.fx_error)
            return False
        self.fx_error = None
        return True

    async def _refresh_prices(self, http: httpx.AsyncClient) -> bool:
        try:
            page = await http.get(BULLETIN_PAGE, timeout=60)
            page.raise_for_status()
            url = find_xlsx_url(page.text)
            xlsx = await http.get(url, timeout=60)
            xlsx.raise_for_status()
            prices, as_of = parse_bulletin(xlsx.content)
        except (httpx.HTTPError, ValueError, KeyError, OSError) as exc:
            self.last_error = f"{exc.__class__.__name__}: {exc}"[:300]
            logger.warning("Base price refresh failed: %s", self.last_error)
            return False
        pm.apply_base_prices(prices, as_of)
        self.last_success, self.last_error, self.source_url = datetime.now(_tz()), None, url
        logger.info("Base prices refreshed: %d countries, bulletin of %s", len(prices), as_of)
        return True

    def is_stale(self, now: datetime | None = None) -> bool:
        now = now or datetime.now(_tz())
        due = last_scheduled_run(now)
        if self.last_success and self.last_success >= due:
            return False
        return not (self.last_attempt and now - self.last_attempt < RETRY_AFTER_FAILURE)

    def ensure_fresh(self, http: httpx.AsyncClient) -> None:
        """Start a background catch-up refresh if the scheduled one was missed."""
        if not self.enabled or not self.is_stale():
            return
        if self._lock.locked() or (self._catch_up and not self._catch_up.done()):
            return
        self._catch_up = asyncio.create_task(self.refresh(http))

    async def _loop(self, http: httpx.AsyncClient) -> None:
        await self.refresh(http)
        while True:
            delay = (next_scheduled_run(datetime.now(_tz())) - datetime.now(_tz())).total_seconds()
            await asyncio.sleep(max(delay, 1))
            await self.refresh(http)

    def start(self, http: httpx.AsyncClient) -> None:
        if self.enabled and self._task is None:
            self._task = asyncio.create_task(self._loop(http))

    async def stop(self) -> None:
        for task in (self._task, self._catch_up):
            if task and not task.done():
                task.cancel()
        self._task = self._catch_up = None

    def status(self) -> dict:
        now = datetime.now(_tz())
        return {
            "auto_refresh": self.enabled,
            "as_of": pm.BASE_PRICES_AS_OF,
            "source": "European Commission Weekly Oil Bulletin (prices with taxes)",
            "source_url": self.source_url,
            "schedule": f"daily {_refresh_time().strftime('%H:%M')} {_tz().key}",
            "next_refresh": next_scheduled_run(now).isoformat(timespec="minutes") if self.enabled else None,
            "last_success": self.last_success.isoformat(timespec="seconds") if self.last_success else None,
            "last_error": self.last_error,
            "fx_as_of": fx.as_of,
            "fx_source": fx.source,
            "fx_error": self.fx_error,
        }


updater = BasePriceUpdater()
